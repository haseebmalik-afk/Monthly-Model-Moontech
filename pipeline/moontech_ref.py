"""
MoonTech dashboard - REFERENCE PIPELINE (Level 0 + Level 1)
Built only from SPEC_dashboard-data.md and config.yaml, to validate the spec and check figures.
It is a reference for the real build, not the production dashboard.

Month index k: 0 = Jan-24 ... 31 = Aug-26 ... 33 = Oct-26.
Run:  python moontech_ref.py <kit_dir> [--merchant-source system|model]
  system (default) = spec-compliant ingestion (S2 history + S1 Datos extracts; system Ounass Jul-26)
  model            = Model's Revenue By Merchant tab for all months (includes the manual Ounass Jul-26 figure)
"""
import sys, re, csv, json, difflib, datetime, os
import openpyxl, yaml

NM = 34
def mo(k): return '%d-%02d' % (2024 + k // 12, k % 12 + 1)
def kidx(y, m): return (y - 2024) * 12 + m - 1

# ----------------------------------------------------------------------------- config
def load_config(kit):
    with open(os.path.join(kit, 'config.yaml')) as f:
        return yaml.safe_load(f)

# ----------------------------------------------------------------------------- name matching (§5)
def clean_name(s, suffixes):
    s = str(s).lower()
    for suf in sorted(suffixes, key=len, reverse=True):
        suf = suf.lower()
        if suf.startswith('('):
            s = s.replace(suf, ' ')
        else:
            s = re.sub(r'\b' + re.escape(suf) + r'\b', ' ', s)
    return re.sub(r'[^a-z0-9]', '', s)

def match_name(name, candidates, cfg):
    """Return (match_type, candidate, score). Order: exact -> contains (unique, >=min len) -> fuzzy -> none."""
    nm = cfg['name_matching']; sx = nm['strip_suffixes']
    n = clean_name(name, sx)
    cl = {c: clean_name(c, sx) for c in candidates}
    ex = [c for c, v in cl.items() if v == n]
    if ex: return 'exact', ex[0], 1.0
    co = [c for c, v in cl.items() if len(v) >= nm['contains_min_length'] and (v in n or n in v)]
    if len(co) == 1: return 'contains', co[0], 1.0
    if not candidates: return 'none', None, 0.0
    best = max(candidates, key=lambda c: difflib.SequenceMatcher(None, n, cl[c]).ratio())
    sc = difflib.SequenceMatcher(None, n, cl[best]).ratio()
    return ('fuzzy' if sc >= nm['fuzzy_min_similarity'] else 'none'), best, round(sc, 2)

# ----------------------------------------------------------------------------- loaders (Level 0)
def load_master(kit):
    rows = list(csv.DictReader(open(os.path.join(kit, 'merchant_master_seed.csv'))))
    master = {}
    for r in rows:
        master[r['merchant']] = {
            'aliases': [a.strip() for a in r['extract_aliases'].split(';') if a.strip()],
            'source': r['source'], 'roas': r['roas_ladder_name'] or None,
            'size_class': r['company_size_class (Growth/SME; Enterprise is computed)'] or 'Unassigned',
            'status': r['status']}
    return master

def load_ladder(kit, cfg):
    path = os.path.join(kit, 'data', 'moontech-roas-ladder-2026-10-05.csv')
    rows = list(csv.reader(open(path, encoding='utf-8-sig')))
    hdr = rows[0]
    mcols = [i for i, h in enumerate(hdr) if '(USD)' in h and 'Total' not in h and 'ARPA' not in h]
    # month parsing: "Nov 25 (USD)", "Dec 25 (USD)", "Jan (USD)" ... (no year => 2026)
    mon = {m: i + 1 for i, m in enumerate(['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec'])}
    colk = {}
    for i in mcols:
        t = hdr[i].replace('(USD)', '').split()
        y = 2000 + int(t[1]) if len(t) > 1 else 2026
        colk[i] = kidx(y, mon[t[0][:3]])
    lvl_re = re.compile(cfg['roas_status']['other_pricing_levels_not_matching'])
    lad = {}
    for r in rows[1:]:
        if len(r) < len(hdr) - 2 or not r[0] or r[0].startswith(('Fleet', 'As of')): continue
        c = dict(zip(hdr, r))
        y, m = map(int, c['Since'].split('-'))
        pay = [0.0] * NM
        for i, k in colk.items():
            if r[i]: pay[k] = float(r[i])
        ch = c.get('Churn month', '')
        churn = kidx(int(ch.split()[1]), mon[ch.split()[0][:3]]) if ch else None
        lad[c['Client']] = {'since': kidx(y, m), 'pay': pay, 'churn': churn, 'level': c['Level'],
                            'other_pricing': not lvl_re.match(c['Level'])}
    return lad

def load_merchant_revenue(kit, master, lad, cfg, source='system'):
    """AED per merchant per month (k 0..31). Returns rev, adjustment_row, review_list."""
    D = os.path.join(kit, 'data')
    model = openpyxl.load_workbook(os.path.join(D, 'MoonTech_Monthly_Model_6.xlsx'), data_only=True)
    rbm = model['Revenue By Merchant  Month']
    rev = {}; adj = [0.0] * 32; review = []
    for r in range(2, 285):
        n = rbm.cell(r, 2).value
        if not n: continue
        vals = [(rbm.cell(r, c).value or 0) for c in range(6, 38)]
        vals = [v if isinstance(v, (int, float)) else 0 for v in vals]
        if n == 'Adjustment': adj = vals; continue
        last = 32 if source == 'model' else 21          # S2 history = Jan-24..Sep-25
        rev[n] = [vals[k] if k < last else 0.0 for k in range(32)]
    if source == 'system':
        up = openpyxl.load_workbook(os.path.join(D, 'Moontech-_Data_updates_2025_1.xlsx'), data_only=True)
        ex = up['Revenue by Merchant']
        alias = {}
        for m, v in master.items():
            alias[m] = m
            for a in v['aliases']: alias[a] = m
        blocks = [(1, [2, 3, 4], [21, 22, 23]), (8, list(range(9, 17)), list(range(24, 32)))]
        for ncol, vcols, ks in blocks:
            for r in range(2, ex.max_row + 1):
                n = ex.cell(r, ncol).value
                if not n or str(n).strip() in ('Adjustment', 'PnL Total') or 'total' in str(n).lower(): continue
                n = str(n).strip()
                if n in alias: m = alias[n]
                else:
                    t, cand, sc = match_name(n, [x for x in master if master[x]['source'] == 'Datos'], cfg)
                    if t in ('exact', 'contains'): m = cand
                    else:
                        t2, cand2, sc2 = match_name(n, list(lad), cfg)
                        m = n
                        review.append({'extract_name': n, 'merchant_match': (t, cand, sc), 'ladder_match': (t2, cand2, sc2),
                                       'action': 'Proposed: new merchant' + (' linked to ' + cand2 if t2 != 'none' else ' (legacy)')})
                        master.setdefault(n, {'aliases': [], 'source': 'Datos', 'roas': cand2 if t2 != 'none' else None,
                                              'size_class': 'Unassigned', 'status': 'Proposed'})
                    alias[n] = m
                rev.setdefault(m, [0.0] * 32)
                for c, k in zip(vcols, ks):
                    v = ex.cell(r, c).value
                    if isinstance(v, (int, float)): rev[m][k] += v
    return rev, adj, review

def load_pl(kit, cfg):
    """P&L lines in AED, k 0..31. 2026 from '2026 Financials'; 2024-25 from Monthly Actuals (USD x FX)."""
    fx = cfg['fx']['aed_per_usd']
    model = openpyxl.load_workbook(os.path.join(kit, 'data', 'MoonTech_Monthly_Model_6.xlsx'), data_only=True)
    fin = model['2026 Financials']; ma = model['Monthly Actuals']
    lines = {'sales': [2], 'creator_payouts': [7, 8, 9], 'materials': [11], 'other_cos': [12, 13], 'freelancers': [14],
             'bd_salaries': [24], 'digital_marketing': [25], 'marketing_events': [26], 'account_mgmt': [27],
             'c_suite': [30], 'ops_salaries': [31], 'advisory': [32], 'development': [35], 'system_expenses': [36],
             'ga': [38], 'da': [44], 'finance_cost': [48], 'tax': [49], 'transfer_pricing': [56], 'pl_gross_profit': [18],
             'pl_ebitda': [42], 'pl_total_cos': [6]}
    ma_rows = {'creator_payouts': [26], 'materials': [27], 'freelancers': [28], 'bd_salaries': [35], 'digital_marketing': [36],
               'marketing_events': [37], 'account_mgmt': [38], 'c_suite': [42], 'ops_salaries': [43], 'advisory': [44],
               'development': [48, 64], 'system_expenses': [49], 'ga': [52], 'da': [57], 'finance_cost': [60], 'tax': [61]}
    pl = {l: [0.0] * 32 for l in lines}
    for l, rs in lines.items():
        for j in range(8):
            pl[l][24 + j] = sum((fin.cell(r, 2 + j).value or 0) for r in rs)
    for l, rs in ma_rows.items():
        for k in range(24):
            pl[l][k] = sum((ma.cell(r, 6 + k).value or 0) for r in rs) * fx
    return pl

def load_creators(kit, cfg):
    """Net orders per Contact ID (text) per month; Unattributed separate. History from Model; Jan-Aug 26 from staging extracts."""
    D = os.path.join(kit, 'data')
    model = openpyxl.load_workbook(os.path.join(D, 'MoonTech_Monthly_Model_6.xlsx'), data_only=True)
    ob = model['Orders by Creator  Month']
    orders = {}; unatt = [0.0] * 32; flags = []
    valid = set(cfg['creator_id_checks']['valid_lengths'])
    def norm_id(v):
        if v in (None, '', 0, '0'): return None
        s = str(v).strip()
        if re.fullmatch(r'\d+\.0', s): s = s[:-2]
        return s
    for r in range(2, ob.max_row + 1):
        i = norm_id(ob.cell(r, 1).value)
        vals = [(ob.cell(r, c).value or 0) for c in range(4, 28)]   # Jan-24..Dec-25 (k 0..23)
        if i is None:
            if ob.cell(r, 1).value == 0:
                for k in range(24): unatt[k] += vals[k]
            continue
        orders.setdefault(i, [0.0] * 32)
        for k in range(24): orders[i][k] += vals[k]
    up = openpyxl.load_workbook(os.path.join(D, 'Moontech-_Data_updates_2025_1.xlsx'), data_only=True)
    st = up['Orders by Creator']
    id_cols = [2, 7, 12, 17, 22, 27, 32, 37]
    for b, ic in enumerate(id_cols):
        k = 24 + b
        for r in range(2, st.max_row + 1):
            v = st.cell(r, ic + 2).value
            if not isinstance(v, (int, float)): continue
            i = norm_id(st.cell(r, ic).value)
            if i is None: unatt[k] += v; continue
            if not i.isdigit() or len(i) not in valid or 'E' in i:
                flags.append(('rejected_format', i, mo(k))); continue
            if i not in orders:
                # rounding-match check (corrupted long IDs)
                for ex in list(orders):
                    if len(ex) >= 10 and any(int(i) == round(int(ex), -(len(ex) - s)) for s in cfg['creator_id_checks']['rounding_match_sig_figs'] if len(ex) > s):
                        flags.append(('rounding_match_reassigned', i, ex)); i = ex; break
                orders.setdefault(i, [0.0] * 32)
            orders[i][k] += round(v + 1e-9)
    return orders, unatt, flags

def load_cac(kit, cfg):
    up = openpyxl.load_workbook(os.path.join(kit, 'data', 'Moontech-_Data_updates_2025_1.xlsx'), data_only=True)
    c26 = up['CAC- 2026']; c = up['CAC']
    merch = [None] * NM; creat = [None] * NM
    for col in range(4, 27):                                  # CAC sheet: Jan-24.. (row 11 merchant USD, row 24 creator USD)
        d = c.cell(2, col).value
        if isinstance(d, datetime.datetime):
            k = kidx(d.year, d.month)
            if k <= kidx(2025, 9):
                merch[k] = c.cell(11, col).value; creat[k] = c.cell(24, col).value
    for i in range(11):                                       # CAC- 2026: Oct-25..Aug-26 (row 14 / 21 USD) - wins on overlap
        k = kidx(2025, 10) + i
        merch[k] = c26.cell(14, 2 + i).value; creat[k] = c26.cell(21, 2 + i).value
    return merch, creat

# ----------------------------------------------------------------------------- Level 1: status rules
ACT = ('New', 'Existing', 'Reactivated')

def legacy_status(x_aed, cfg, stop=None):
    fx = cfg['fx']['aed_per_usd']; ms = cfg['merchant_status']
    T = ms['active_min_annualised_usd'] * fx; Q = ms['quiet_month_threshold_usd'] * fx; NQ = ms['quiet_months_to_churn']
    out = ['-'] * NM; start = None; state = False; churned_before = False
    for k in range(32):
        if stop is not None and k >= stop: break
        if start is None and x_aed[k] > 1: start = k
        if start is None: continue
        ten = k - start + 1
        val = sum(x_aed[max(start, k - 11):k + 1]) if ten >= 12 else sum(x_aed[start:k + 1]) * 12 / ten
        quiet = k - NQ + 1 >= 0 and all(a < Q for a in x_aed[k - NQ + 1:k + 1])
        act = val >= T and not quiet
        if act and not state: lab = 'Reactivated' if churned_before else 'New'
        elif act: lab = 'Existing'
        elif state: lab = 'Churned'; churned_before = True; start = None
        else: lab = '-'
        state = act; out[k] = lab
    return out

def roas_status(b, cfg):
    since, pay, churn = b['since'], b['pay'], b['churn']
    if churn is not None and churn <= since: churn = since + 1           # Since-month rule
    NR = cfg['roas_status']['months_without_payment_to_churn']
    out = ['-'] * NM; state = False; churned = False; last = None
    for k in range(since, NM):
        if pay[k] > 0: last = k
        if churned:
            act = pay[k] > 0                                              # R3 reactivation
            if act: churn = None
        elif churn is not None and k >= churn: act = False               # R1
        else:
            ref = last if last is not None else since
            act = (k - ref) < NR or k == since                           # R2 backstop; always active in Since month
        if act and not state: out[k] = 'Reactivated' if churned else 'New'
        elif act: out[k] = 'Existing'
        elif state: out[k] = 'Churned'; churned = True
        state = act
    return out

def creator_status(x, cfg):
    cs = cfg['creator_status']; thr = cs['active_min_avg_net_orders_per_month']; W = cs['activity_window_months']
    out = ['-'] * 32; start = None; state = False; churned = False
    for k in range(32):
        if start is None and x[k] >= 1: start = k
        recent = any(v >= 1 for v in x[max(0, k - W + 1):k + 1])
        act = False
        if start is not None:
            w = min(12, k - start + 1)
            act = sum(x[k - w + 1:k + 1]) / w >= thr and recent
        if act and not state: lab = 'Reactivated' if churned else 'New'
        elif act: lab = 'Existing'
        elif state: lab = 'Churned'; churned = True; start = None
        elif recent: lab = 'Emerging'
        else: lab = '-'
        state = act; out[k] = lab
    return out

def annualised(x, k):
    f = next((j for j, a in enumerate(x) if a > 1), None)
    if f is None or k < f: return 0.0
    t = k - f + 1
    return sum(x[max(f, k - 11):k + 1]) if t >= 12 else sum(x[f:k + 1]) * 12 / t

def hysteresis(x, enter, exit_, n=32):
    st = False; out = []
    for k in range(n):
        a = annualised(x, k)
        if not st and a >= enter: st = True
        elif st and a < exit_: st = False
        out.append(st)
    return out

# ----------------------------------------------------------------------------- engine
class Engine:
    def __init__(self, kit, source='system'):
        self.cfg = cfg = load_config(kit); self.fx = cfg['fx']['aed_per_usd']
        self.master = load_master(kit)
        self.lad = load_ladder(kit, cfg)
        self.rev, self.adj, self.review = load_merchant_revenue(kit, self.master, self.lad, cfg, source)
        self.pl = load_pl(kit, cfg)
        self.orders, self.unatt, self.id_flags = load_creators(kit, cfg)
        self.cac_m, self.cac_c = load_cac(kit, cfg)
        fx = self.fx
        # ROAS links
        self.link = {m: v['roas'] for m, v in self.master.items() if v['roas'] and m in self.rev}
        self.since = {m: self.lad[b]['since'] for m, b in self.link.items()}
        # statuses
        self.leg = {m: legacy_status(x, cfg, self.since.get(m)) for m, x in self.rev.items()}
        self.leg_full = {m: legacy_status(x, cfg) for m, x in self.rev.items()}
        self.rst = {b: roas_status(v, cfg) for b, v in self.lad.items()}
        self.cst = {i: creator_status(x, cfg) for i, x in self.orders.items()}
        # revenue composition (USD)
        self.roas = [sum(v['pay'][k] for v in self.lad.values() if not v['other_pricing']) for k in range(NM)]
        self.other = [sum(v['pay'][k] for v in self.lad.values() if v['other_pricing']) for k in range(NM)]
        self.named = [sum(x[k] for m, x in self.rev.items() if m not in self.since or k < self.since[m]) / fx for k in range(32)]
        self.total = [(self.pl['sales'][k] / fx) if k >= 24 else (sum(x[k] for x in self.rev.values()) + self.adj[k]) / fx for k in range(32)]
        self.plug = [self.total[k] - self.named[k] - self.roas[k] - self.other[k] for k in range(32)]
        self.legacy = [self.total[k] - self.roas[k] - self.other[k] for k in range(32)]
        # seasonality weights (rule): avg monthly share of legacy revenue over last 2 full years (2024-25)
        tot = sum(self.legacy[:24])
        self.w = [(self.legacy[i] + self.legacy[12 + i]) / tot for i in range(12)]
        # segments (Enterprise annualised hysteresis on reported revenue: Datos pre-Since + ladder post-Since)
        sg = cfg['segmentation']
        self.seg = {}
        for m, x in self.rev.items():
            rep = list(x)
            if m in self.since:
                b = self.link[m]
                rep = [x[k] if k < self.since[m] else self.lad[b]['pay'][k] * fx for k in range(32)]
            ent = hysteresis(rep, sg['enterprise_enter_annualised_aed'], sg['enterprise_exit_annualised_aed'])
            base = self.master.get(m, {}).get('size_class', 'Unassigned')
            self.seg[m] = ['Enterprise' if ent[k] else base for k in range(32)]
        self.major = {i: hysteresis(x, sg['major_creator_enter_annualised_orders'], sg['major_creator_exit_annualised_orders'])
                      for i, x in self.orders.items()}

    # --- helpers
    def legacy_active(self, k, seg=None):
        return [m for m in self.rev if self.leg[m][k] in ACT and (seg is None or self.seg[m][k] in seg)]
    def roas_active(self, k, include_other=False):
        return [b for b, v in self.lad.items() if self.rst[b][k] in ACT and (include_other or not v['other_pricing'])]
    def roas_arr(self, k): return sum(self.roas[j] for j in range(k - 2, k + 1)) * self.cfg['arr']['roas_multiplier']
    def legacy_runrate(self, k): return sum(self.legacy[j] / self.w[j % 12] for j in range(k - 2, k + 1)) / 3
    def seg_arpa(self, k, seg):
        R = N = 0
        for j in range(k - 2, k + 1):
            act = self.legacy_active(j, seg)
            R += sum(self.rev[m][j] for m in act) / self.fx / self.w[j % 12] / 3
            N += len(act) / 3
        return R / N, N
    def roas_arpa(self, k):
        N = sum(len(self.roas_active(j)) for j in range(k - 2, k + 1)) / 3
        return self.roas_arr(k) / N, N
    def nrr(self, series, p0, p1):
        s = e = keep = 0
        for v in series.values():
            a = sum(v[j] for j in p0); b = sum(v[j] for j in p1)
            if a > 0: s += a; e += b; keep += min(a, b)
        return e / s, keep / s
    def legacy_series(self, segfilter=None, at=None):
        out = {}
        for m, x in self.rev.items():
            if segfilter and self.seg[m][at] not in segfilter: continue
            out[m] = [x[k] / self.fx if (m not in self.since or k < self.since[m]) else 0.0 for k in range(32)]
        return out
    def roas_series(self): return {b: v['pay'] for b, v in self.lad.items() if not v['other_pricing']}

    # --- check figures
    def compute_checks(self):
        r = {}; fx = self.fx; pl = self.pl
        J, A, Ju, Au = kidx(2026, 1), kidx(2026, 4), kidx(2026, 7), kidx(2026, 8)
        r['CF01'] = self.total[J] / 1e3; r['CF02'] = self.named[J] / 1e3; r['CF03'] = self.roas[J] / 1e3
        r['CF04'] = self.plug[J] / 1e3; r['CF05'] = self.plug[A] / 1e3; r['CF06'] = self.plug[Ju] / 1e3; r['CF07'] = self.plug[Au] / 1e3
        ytd = range(24, 32)
        S = lambda l: sum(pl[l][k] for k in ytd) / fx
        rev = S('sales'); cp = S('creator_payouts'); cos = cp + S('materials') + S('other_cos') + S('freelancers')
        opex = sum(S(l) for l in ['bd_salaries', 'digital_marketing', 'marketing_events', 'account_mgmt', 'c_suite', 'ops_salaries',
                                  'advisory', 'development', 'system_expenses', 'ga'])
        r['CF08'] = rev / 1e3; r['CF09'] = cp / 1e3; r['CF10'] = (1 - cp / rev) * 100; r['CF11'] = (1 - cos / rev) * 100
        r['CF12'] = (rev - cos - opex) / 1e3
        r['_check_cos_vs_pl'] = (cos - S('pl_total_cos')) ; r['_check_ebitda_vs_pl'] = (rev - cos - opex - S('pl_ebitda'))
        r['CF13'] = sum(1 for m in self.rev if self.leg_full[m][Au] in ACT)
        r['CF14'] = len(self.legacy_active(Au)); r['CF15'] = len(self.roas_active(Au, include_other=True))
        r['CF16'] = r['CF14'] + r['CF15']
        sep = kidx(2026, 9)
        r['CF17'] = sum(1 for v in self.lad.values() if not v['other_pricing'] and v['since'] <= sep)
        r['CF18'] = len(self.roas_active(sep))
        act_c = [i for i in self.orders if self.cst[i][Au] in ACT]
        r['CF19'] = len(act_c); r['CF20'] = sum(1 for i in self.orders if self.cst[i][Au] == 'Emerging')
        r['CF21'] = sum(1 for i in self.orders if self.major[i][Au])
        r['CF22'] = sum(x[Au] for x in self.orders.values()) + self.unatt[Au]
        r['CF23'] = self.roas_arr(sep)
        r['CF24'] = self.legacy_runrate(Ju) / 1e6; r['CF25'] = self.legacy_runrate(Au) / 1e6
        r['CF26'] = self.seg_arpa(Ju, ['Enterprise'])[0] / 1e6
        r['CF27'] = self.seg_arpa(Ju, ['Growth'])[0] / 1e3
        r['CF28'] = self.roas_arpa(sep)[0]
        ttm = range(kidx(2025, 9), Au + 1)
        r['CF29'] = len({m for m in self.rev for k in ttm if self.leg[m][k] == 'Churned'})
        r['CF30'] = sum(len(self.legacy_active(k)) for k in ttm) / 12
        launch = kidx(2025, 11); hist = range(launch, sep + 1)
        r['CF31'] = len({b for b, v in self.lad.items() if not v['other_pricing'] for k in hist if self.rst[b][k] == 'Churned'})
        L = self.legacy_series()
        n12, g12 = self.nrr(L, range(kidx(2024, 9), kidx(2025, 8) + 1), ttm)
        r['CF32'] = n12 * 100; r['CF33'] = g12 * 100
        R = self.roas_series()
        r['CF34'] = self.nrr(R, range(24, 27), range(27, 30))[0] * 100
        r['CF35'] = self.nrr(R, range(27, 30), range(30, 33))[0] * 100
        spend = lambda a, b: sum((self.cac_m[k] or 0) for k in range(a, b + 1))
        newb = lambda a, b: sum(1 for v in self.lad.values() if not v['other_pricing'] and a <= v['since'] <= b)
        r['CF36'] = (spend(launch, Au) + 8500) / 37
        r['CF37'] = spend(27, 29) / newb(27, 29)
        arpa_q3 = r['CF28']
        churned = r['CF31']; avg_act = sum(len(self.roas_active(k)) for k in hist) / len(hist)
        m_ch = churned / avg_act / len(hist)
        life = min(self.cfg['ltv']['lifetime_cap_months'], 1 / m_ch)
        r['CF38'] = arpa_q3 / 12 * self.cfg['ltv']['roas_margin'] * life
        r['_roas_monthly_churn'] = m_ch
        # legacy Growth & SME LTV (revenue churn basis)
        rev_ttm = sum(self.total[k] for k in ttm); cp_ttm = sum(pl['creator_payouts'][k] for k in ttm) / fx
        margin = 1 - cp_ttm / rev_ttm
        r['_legacy_margin_ttm'] = margin
        arpa_g, _ = self.seg_arpa(Ju, ['Growth', 'SME'])
        G = self.legacy_series(segfilter=['Growth', 'SME', 'Unassigned'], at=kidx(2025, 8))
        _, grr_g = self.nrr(G, range(kidx(2024, 9), kidx(2025, 8) + 1), ttm)
        mg = (1 - grr_g) / 12
        lifeg = min(self.cfg['ltv']['lifetime_cap_months'], 1 / mg) if mg > 0 else 60
        r['CF39'] = arpa_g / 12 * margin * lifeg / 1e3
        r['_growth_grr'] = grr_g
        # concentration (accounts = reported revenue: Datos pre-Since + ladder post-Since)
        acc = {}
        for m, x in self.rev.items():
            acc[m] = sum((x[k] / fx if (m not in self.since or k < self.since[m]) else 0) for k in ttm)
            if m in self.link: acc[m] += sum(self.lad[self.link[m]]['pay'][k] for k in ttm if k >= self.since[m])
        for b, v in self.lad.items():
            if b not in self.link.values(): acc['ROAS:' + b] = sum(v['pay'][k] for k in ttm)
        top = sorted(acc.values(), reverse=True); T = sum(self.total[k] for k in ttm)
        r['CF40'] = top[0] / T * 100; r['CF41'] = sum(top[:5]) / T * 100
        r['CF42'] = (sum(self.total[24:32]) / sum(self.total[12:20]) - 1) * 100
        legytd = sum(self.legacy[24:32]); share = sum(self.w[:8])
        roas_ytd = sum(self.roas[24:32]) + sum(self.other[24:32])
        r['CF43'] = (legytd / share + roas_ytd + self.roas_arr(Au) * 4 / 12) / 1e6
        r['_CF43_parts'] = {'legacy': legytd / share / 1e6, 'roas': (roas_ytd + self.roas_arr(Au) * 4 / 12) / 1e6,
                            'roas_simple_x12_8': roas_ytd * 12 / 8 / 1e6}
        return r

if __name__ == '__main__':
    kit = sys.argv[1] if len(sys.argv) > 1 else '.'
    src = 'model' if '--merchant-source' in sys.argv and sys.argv[sys.argv.index('--merchant-source') + 1] == 'model' else 'system'
    E = Engine(kit, src)
    res = E.compute_checks()
    exp = list(csv.DictReader(open(os.path.join(kit, 'check_figures.csv'))))
    out = []
    col = 'expected_model_dataset' if src == 'model' else 'expected_spec_ingestion'
    for c in exp:
        got = res.get(c['id'])
        ok = got is not None and abs(got - float(c[col])) <= float(c['tolerance']) + 1e-9
        out.append({'id': c['id'], 'metric': c['metric'], 'period': c['period'], 'expected': float(c[col]),
                    'got': round(got, 3) if got is not None else None, 'pass': ok})
    print(json.dumps({'merchant_source': src, 'results': out,
                      'diagnostics': {k: v for k, v in res.items() if k.startswith('_')},
                      'review_list': E.review, 'creator_id_flags': E.id_flags[:20],
                      'n_merchants': len(E.rev), 'n_creators': len(E.orders),
                      'weights': [round(x, 4) for x in E.w]}, indent=1, default=str))
